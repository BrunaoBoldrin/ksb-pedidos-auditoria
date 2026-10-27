import unittest
from html.parser import HTMLParser
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import ui_components

from ui_components import (
    cabecalho_item_html,
    filtrar_itens,
    painel_diagnostico_html,
    resumo_lote,
)


class MarkupReader(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags = []
        self.text = []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)

    def handle_data(self, text):
        self.text.append(text)


class PresentationTests(unittest.TestCase):
    def setUp(self):
        self.frame = pd.DataFrame([
            {"Pedido": "4501", "Item": "00010", "Código Material": "02154693", "Status": "OK", "Status Comercial": "OK"},
            {"Pedido": "4501", "Item": "00020", "Código Material": "WGS-5520", "Status": "OK", "Status Comercial": "PENDENTE - REVISÃO DE PREÇO", "Preço Pedido KSB": 100, "Preço Cadastrado": 100},
            {"Pedido": "4502", "Item": "00100", "Código Material": "00123456", "Status": "DIVERGENTE", "Status Comercial": "OK"},
            {"Pedido": "4502", "Item": "00110", "Código Material": "ABC-1", "Status": None, "Status Comercial": pd.NA},
        ])

    def test_summary_does_not_approve_expired_equal_prices_or_incomplete_items(self):
        self.assertEqual(resumo_lote(self.frame), {
            "total": 4, "aprovados": 1, "pendentes": 3,
            "revisao_preco": 1, "divergencia_fiscal": 1,
        })

    def test_empty_or_missing_statuses_never_count_as_approved(self):
        self.assertEqual(resumo_lote(pd.DataFrame())["total"], 0)
        result = resumo_lote(pd.DataFrame([{"Item": "00010"}]))
        self.assertEqual((result["aprovados"], result["pendentes"]), (0, 1))

    def test_visual_filter_preserves_text_codes_and_does_not_mutate_source(self):
        original = self.frame.copy(deep=True)
        pending = filtrar_itens(self.frame, "Com pendência")
        self.assertEqual(pending["Item"].tolist(), ["00020", "00100", "00110"])
        self.assertEqual(pending.iloc[0]["Código Material"], "WGS-5520")
        selected = filtrar_itens(self.frame, "Aprovados", "4501")
        self.assertEqual(selected["Item"].tolist(), ["00010"])
        self.assertTrue(filtrar_itens(self.frame, "Todos", "inexistente").empty)
        pd.testing.assert_frame_equal(original, self.frame)

    def test_untrusted_item_content_is_rendered_as_text(self):
        markup = cabecalho_item_html({
            "Pedido": "4501", "Item": "00020", "Código Material": "WGS-5520",
            "Descrição": '<script>alert("x")</script><img src=x onerror=alert(1)>',
            "Status": "OK", "Status Comercial": '<svg onload="attack">',
        })
        reader = MarkupReader()
        reader.feed(markup)
        self.assertNotIn("script", reader.tags)
        self.assertNotIn("img", reader.tags)
        self.assertNotIn("svg", reader.tags)
        self.assertIn('<script>alert("x")</script>', "".join(reader.text))
        self.assertIn("00020", "".join(reader.text))
        self.assertIn("WGS-5520", "".join(reader.text))

    def test_pending_diagnostic_retains_complete_message_and_line_breaks(self):
        message = "PREÇO NECESSITA DE REVISÃO: ATUALIZE O VALOR DO ITEM.\nÚltima revisão: 22/07/2026"
        markup = painel_diagnostico_html("Comercial", "PENDENTE - REVISÃO DE PREÇO", message)
        reader = MarkupReader()
        reader.feed(markup)
        self.assertIn("br", reader.tags)
        self.assertIn("PREÇO NECESSITA DE REVISÃO: ATUALIZE O VALOR DO ITEM.", "".join(reader.text))
        self.assertIn("Última revisão: 22/07/2026", "".join(reader.text))

    def test_excel_export_contains_business_columns_and_preserves_text_codes(self):
        exporter = getattr(ui_components, "dataframe_to_excel_bytes", None)
        self.assertTrue(callable(exporter), "A exportação Excel das tabelas deve estar disponível")
        frame = pd.DataFrame(
            [{"Selecionar": False, "Código Material": "001492536", "Descrição": "PROTETOR ACOPL"}]
        )

        workbook = load_workbook(BytesIO(exporter(frame, "Itens")), read_only=True)
        rows = list(workbook["Itens"].values)

        self.assertEqual(rows[0], ("Código Material", "Descrição"))
        self.assertEqual(rows[1], ("001492536", "PROTETOR ACOPL"))

    def test_brazilian_currency_format_uses_symbol_thousands_and_decimal_comma(self):
        formatter = getattr(ui_components, "formatar_moeda_brl", None)
        self.assertTrue(callable(formatter), "A formatação monetária brasileira deve estar disponível")
        self.assertEqual(formatter(1234.5), "R$ 1.234,50")
        self.assertEqual(formatter(-98.7), "R$ -98,70")
        self.assertEqual(formatter("1.234,56"), "R$ 1.234,56")
        self.assertEqual(formatter(None), "")

    def test_currency_columns_are_formatted_without_mutating_business_data(self):
        frame = pd.DataFrame([
            {
                "Código Material": "001492536",
                "Preço Unitário Líquido": 1234.5,
                "Preço Anterior": 1000,
                "Quantidade": 2,
            }
        ])

        format_columns = getattr(ui_components, "format_currency_columns", None)
        self.assertTrue(callable(format_columns), "As tabelas devem formatar suas colunas monetárias")
        displayed = format_columns(frame)

        self.assertEqual(displayed.loc[0, "Preço Unitário Líquido"], "R$ 1.234,50")
        self.assertEqual(displayed.loc[0, "Preço Anterior"], "R$ 1.000,00")
        self.assertEqual(displayed.loc[0, "Quantidade"], 2)
        self.assertEqual(frame.loc[0, "Preço Unitário Líquido"], 1234.5)

    def test_excel_keeps_price_numeric_and_applies_brazilian_currency_style(self):
        frame = pd.DataFrame([
            {"Código Material": "001492536", "Preço Unitário Líquido": 1234.5}
        ])

        workbook = load_workbook(BytesIO(ui_components.dataframe_to_excel_bytes(frame, "Itens")))
        price_cell = workbook["Itens"]["B2"]

        self.assertEqual(price_cell.value, 1234.5)
        self.assertEqual(price_cell.number_format, '[$R$-pt-BR] #,##0.00')


if __name__ == "__main__":
    unittest.main()
